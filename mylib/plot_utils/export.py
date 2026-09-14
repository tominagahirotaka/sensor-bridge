"""Header export: writes a recording's Header object (subject/board/time
metadata + checkpoint marks) out to an .xlsx workbook."""

import time

import pandas as pd
from openpyxl import load_workbook

from mylib.i18n import t as tr


def save_header(header, excelfile):
    start = time.perf_counter()
    column_cpt = ['real time', 'time']
    with pd.ExcelWriter(excelfile) as writer:
        pd.DataFrame(header.time['check_points'],
                     columns=column_cpt,
                     index=range(1,
                                 len(header.time['check_points']) + 1)).to_excel(writer,
                                                                                 sheet_name='Mark_points',
                                                                                 index=True,
                                                                                 header=True)
    wb = load_workbook(filename=excelfile)
    for sheet, dic in header.__dict__.items():
        ws = wb.create_sheet(sheet)
        if sheet == 'machine' and type(dic) is not dict:
            ws.append([str(dic)])
        else:
            for key, lst in dic.items():
                ws.append([key, str(lst)])
    wb.save(excelfile)
    wb.close()
    print(tr('pu.auto_saving_header', excelfile=excelfile, sec=time.perf_counter() - start))
